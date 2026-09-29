import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='003366')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            b_el = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val","single")}" w:sz="{edge_data.get("sz",4)}" w:space="0" w:color="{edge_data.get("color","auto")}"/>')
            tcBorders.append(b_el)
        else:
            b_el = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_el)
    tcPr.append(tcBorders)

def add_callout(doc, text_lines, title="PUNTO CLAVE / REGLA CLÍNICA", border_color="0D7685", bg_color="F0F9FF"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    set_cell_border(cell, left=dict(sz=24, val='single', color=border_color),
                          top=dict(val='none'), bottom=dict(val='none'), right=dict(val='none'))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(10.5)
    run_t.font.bold = True
    r, g, b = int(border_color[0:2], 16), int(border_color[2:4], 16), int(border_color[4:6], 16)
    run_t.font.color.rgb = RGBColor(r, g, b)
    
    for line in text_lines:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.line_spacing = 1.15
        run = p2.add_run(line)
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(4)
    p_spacer.paragraph_format.space_after = Pt(4)

def parse_and_build_docx(md_path, docx_path, doc_title, doc_badge, primary_color_hex="1B365D", secondary_color_hex="0D7685"):
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("UNSTA — Cátedra de Nutrición Infantil | Material Docente Oficial")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Página ")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        f_field = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        fp._p.append(f_field)

    # Document Header Box / Banner
    banner_tbl = doc.add_table(rows=1, cols=1)
    banner_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner_tbl.autofit = False
    bcell = banner_tbl.cell(0, 0)
    bcell.width = Inches(6.8)
    set_cell_background(bcell, "F8FAFC")
    set_cell_margins(bcell, top=180, bottom=180, left=200, right=200)
    set_cell_border(bcell, left=dict(sz=28, val='single', color=primary_color_hex),
                           top=dict(sz=4, val='single', color="E2E8F0"),
                           bottom=dict(sz=12, val='single', color=primary_color_hex),
                           right=dict(sz=4, val='single', color="E2E8F0"))
    
    bp = bcell.paragraphs[0]
    bp.paragraph_format.space_before = Pt(0)
    bp.paragraph_format.space_after = Pt(2)
    brun_sub = bp.add_run(f"UNIVERSIDAD DEL NORTE SANTO TOMÁS DE AQUINO (UNSTA) — {doc_badge.upper()}\n")
    brun_sub.font.name = "Calibri"
    brun_sub.font.size = Pt(9)
    brun_sub.font.bold = True
    r, g, b = int(secondary_color_hex[0:2], 16), int(secondary_color_hex[2:4], 16), int(secondary_color_hex[4:6], 16)
    brun_sub.font.color.rgb = RGBColor(r, g, b)
    
    brun_title = bp.add_run(doc_title)
    brun_title.font.name = "Calibri"
    brun_title.font.size = Pt(17)
    brun_title.font.bold = True
    pr, pg, pb = int(primary_color_hex[0:2], 16), int(primary_color_hex[2:4], 16), int(primary_color_hex[4:6], 16)
    brun_title.font.color.rgb = RGBColor(pr, pg, pb)
    
    bp_doc = bcell.add_paragraph()
    bp_doc.paragraph_format.space_before = Pt(4)
    bp_doc.paragraph_format.space_after = Pt(0)
    brun_doc = bp_doc.add_run("Facultad de Ciencias de la Salud • Carrera de Nutrición • Cátedra de Nutrición Infantil\nProfesor a Cargo: Lic. Germán")
    brun_doc.font.name = "Calibri"
    brun_doc.font.size = Pt(9.5)
    brun_doc.font.italic = True
    brun_doc.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    doc.add_paragraph().paragraph_format.space_before = Pt(8)

    # Read Markdown lines
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_lines = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle code/box blocks
        if stripped.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                # Render code box as callout
                if code_lines:
                    title_call = "ESQUEMA SINTÉTICO / PROTOCOLO CLÍNICO"
                    first = code_lines[0]
                    if "PROTOCOLO" in first or "PRINCIPIOS" in first or "PATRONES" in first:
                        title_call = re.sub(r'[^A-Za-z0-9ÁÉÍÓÚáéíóúñÑ\s\(\)]', '', first).strip()
                        code_lines = code_lines[1:]
                    clean_lines = [l.strip('│┌┐└┘├┤─').strip() for l in code_lines if l.strip('│┌┐└┘├┤─').strip()]
                    if clean_lines:
                        add_callout(doc, clean_lines, title=title_call, border_color=secondary_color_hex, bg_color="F8FAFC")
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Skip main H1 since banner handled it
        if stripped.startswith('# '):
            i += 1
            continue

        # Skip horizontal dividers
        if stripped in ('---', '***', '___'):
            i += 1
            continue

        # Headings H2
        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = "Calibri"
            run.font.size = Pt(13.5)
            run.font.bold = True
            pr, pg, pb = int(primary_color_hex[0:2], 16), int(primary_color_hex[2:4], 16), int(primary_color_hex[4:6], 16)
            run.font.color.rgb = RGBColor(pr, pg, pb)
            i += 1
            continue

        # Headings H3
        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(11)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run.font.bold = True
            r, g, b = int(secondary_color_hex[0:2], 16), int(secondary_color_hex[2:4], 16), int(secondary_color_hex[4:6], 16)
            run.font.color.rgb = RGBColor(r, g, b)
            i += 1
            continue

        # Markdown Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Parse table
            if len(table_lines) >= 2:
                header_row = [c.strip() for c in table_lines[0].strip('|').split('|')]
                data_rows = []
                for t_line in table_lines[2:]:  # skip separator
                    data_rows.append([c.strip() for c in t_line.strip('|').split('|')])
                
                tbl = doc.add_table(rows=len(data_rows) + 1, cols=len(header_row))
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = True
                
                # Header row format
                hdr_cells = tbl.rows[0].cells
                for col_idx, h_text in enumerate(header_row):
                    cell = hdr_cells[col_idx]
                    set_cell_background(cell, primary_color_hex)
                    set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                    set_cell_border(cell, top=dict(sz=6, color=primary_color_hex),
                                          bottom=dict(sz=12, color=secondary_color_hex),
                                          left=dict(sz=2, color="E2E8F0"),
                                          right=dict(sz=2, color="E2E8F0"))
                    cp = cell.paragraphs[0]
                    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    clean_h = re.sub(r'[*_`]', '', h_text)
                    crun = cp.add_run(clean_h)
                    crun.font.name = "Calibri"
                    crun.font.size = Pt(9.5)
                    crun.font.bold = True
                    crun.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                
                # Data rows
                for r_idx, r_data in enumerate(data_rows):
                    row_cells = tbl.rows[r_idx + 1].cells
                    bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                    for col_idx, cell_value in enumerate(r_data):
                        if col_idx < len(row_cells):
                            cell = row_cells[col_idx]
                            set_cell_background(cell, bg)
                            set_cell_margins(cell, top=70, bottom=70, left=120, right=120)
                            set_cell_border(cell, top=dict(sz=2, color="E2E8F0"),
                                                  bottom=dict(sz=2, color="E2E8F0"),
                                                  left=dict(sz=2, color="E2E8F0"),
                                                  right=dict(sz=2, color="E2E8F0"))
                            cp = cell.paragraphs[0]
                            cp.paragraph_format.space_before = Pt(0)
                            cp.paragraph_format.space_after = Pt(0)
                            clean_val = re.sub(r'[*_`\$]', '', cell_value).replace('\\text', '').replace('{', '').replace('}', '').strip()
                            crun = cp.add_run(clean_val)
                            crun.font.name = "Calibri"
                            crun.font.size = Pt(9)
                            if "**" in cell_value:
                                crun.font.bold = True
                            crun.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
                
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_before = Pt(4)
                p_spacer.paragraph_format.space_after = Pt(4)
            continue

        # Bullet lists
        if stripped.startswith('* ') or stripped.startswith('- ') or stripped.startswith('+ '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            raw_text = stripped[2:].strip()
            # Process inline bold
            parts = re.split(r'(\*\*.*?\*\*)', raw_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    clean = re.sub(r'[`\$]', '', part).replace('\\text', '').replace('{', '').replace('}', '')
                    run = p.add_run(clean)
                run.font.name = "Calibri"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            i += 1
            continue

        # Numbered lists
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            num = num_match.group(1)
            raw_text = num_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            run_num = p.add_run(f"{num}. ")
            run_num.font.name = "Calibri"
            run_num.font.size = Pt(10)
            run_num.font.bold = True
            run_num.font.color.rgb = RGBColor(pr, pg, pb)
            
            parts = re.split(r'(\*\*.*?\*\*)', raw_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    clean = re.sub(r'[`\$]', '', part).replace('\\text', '').replace('{', '').replace('}', '')
                    run = p.add_run(clean)
                run.font.name = "Calibri"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            i += 1
            continue

        # Blockquote ("> ...")
        if stripped.startswith('> '):
            quote_text = stripped[2:].strip().replace('*', '').replace('"', '')
            add_callout(doc, [f'"{quote_text}"'], title="DECLARACIÓN INSTITUCIONAL OFICIAL", border_color=secondary_color_hex, bg_color="F8FAFC")
            i += 1
            continue

        # Checklist "[ ] ..."
        if stripped.startswith('* [ ]') or stripped.startswith('- [ ]'):
            raw_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            run_box = p.add_run("☐  ")
            run_box.font.name = "Calibri"
            run_box.font.size = Pt(11)
            run_box.font.bold = True
            run_box.font.color.rgb = RGBColor(r, g, b)
            parts = re.split(r'(\*\*.*?\*\*)', raw_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    clean = re.sub(r'[`\$]', '', part).replace('\\text', '').replace('{', '').replace('}', '')
                    run = p.add_run(clean)
                run.font.name = "Calibri"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            i += 1
            continue

        # Standard Paragraph
        if stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            
            # Handle math formulas if standalone
            if stripped.startswith('$$') and stripped.endswith('$$'):
                formula = stripped[2:-2].strip().replace('\\text', '').replace('\\frac', '').replace('{', ' ').replace('}', ' ')
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(formula)
                run.font.name = "Cambria Math"
                run.font.size = Pt(11)
                run.font.bold = True
                run.font.color.rgb = RGBColor(pr, pg, pb)
                i += 1
                continue

            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    clean = re.sub(r'[`\$]', '', part).replace('\\text', '').replace('{', '').replace('}', '')
                    run = p.add_run(clean)
                run.font.name = "Calibri"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        i += 1

    doc.save(docx_path)
    print(f"Successfully generated {docx_path}")

if __name__ == '__main__':
    base_dir = r"c:\Users\germa\OneDrive\Desktop\ARCHIVOS\UNSTA\pagina nutricion infantil\catedra\documentos_teoricos"
    
    # 1. Antropometria
    md_antropo = os.path.join(base_dir, "Manual_Teorico_Practico_Antropometria_UNSTA.md")
    docx_antropo = os.path.join(base_dir, "Manual_Teorico_Practico_Antropometria_UNSTA.docx")
    parse_and_build_docx(md_antropo, docx_antropo, 
                         doc_title="Manual Teórico-Práctico: Evaluación Antropométrica en Pediatría y Gestación", 
                         doc_badge="Documento de Cátedra N° 1 — Antropometría y Composición Corporal",
                         primary_color_hex="1B365D", secondary_color_hex="0D7685")
    
    # 2. Vegetarianismo
    md_veg = os.path.join(base_dir, "Guia_Clinica_Vegetarianismo_Pediatria_UNSTA.md")
    docx_veg = os.path.join(base_dir, "Guia_Clinica_Vegetarianismo_Pediatria_UNSTA.docx")
    parse_and_build_docx(md_veg, docx_veg, 
                         doc_title="Guía Clínica: Planificación y Manejo de Dietas Basadas en Plantas en Pediatría", 
                         doc_badge="Documento de Cátedra N° 2 — Nutrición Basada en Plantas y Vitamina B12",
                         primary_color_hex="164E63", secondary_color_hex="059669")
